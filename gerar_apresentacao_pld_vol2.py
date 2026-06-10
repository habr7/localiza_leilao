# -*- coding: utf-8 -*-
"""
Apresentacao executiva (volume 2):
"Tipologias Especificas de LD/FT em Adquirencia"
Para cada tipologia: descricao, indicadores de alerta, impacto na adquirente,
controles e regras de monitoramento sugeridas.
Tema corporativo azul, alinhado ao deck anterior.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- Paleta ----
NAVY      = RGBColor(0x0F, 0x25, 0x40)
PRIMARY   = RGBColor(0x1F, 0x4E, 0x79)
SECONDARY = RGBColor(0x2E, 0x75, 0xB6)
ACCENT    = RGBColor(0x4A, 0x90, 0xD2)
ICE       = RGBColor(0xEA, 0xF2, 0xFB)
LIGHT     = RGBColor(0xD6, 0xE4, 0xF0)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAYTXT   = RGBColor(0x3A, 0x3A, 0x3A)
RED       = RGBColor(0xC0, 0x39, 0x2B)
GREEN     = RGBColor(0x1E, 0x7A, 0x4D)
AMBER     = RGBColor(0xC9, 0x8A, 0x1B)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None, line_w=None, shape_type=MSO_SHAPE.RECTANGLE):
    shp = s.shapes.add_shape(shape_type, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    return tb, tf


def add_par(tf, text, size=14, color=GRAYTXT, bold=False, italic=False,
            align=PP_ALIGN.LEFT, font=FONT, space_after=4, space_before=0,
            first=False, line_spacing=1.0):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font; f.color.rgb = color
    return p


def bullet(tf, text, size=13, color=GRAYTXT, bold=False, first=False,
           marker="• ", space_after=4, mcolor=None):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.0
    rm = p.add_run(); rm.text = marker
    rm.font.size = Pt(size); rm.font.bold = True; rm.font.name = FONT
    rm.font.color.rgb = mcolor or SECONDARY
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    return p


_pageno = {"n": 0}


def page_header(s, kicker, title, idx=None):
    rect(s, 0, 0, SW, Inches(1.15), PRIMARY)
    rect(s, 0, Inches(1.15), SW, Inches(0.06), ACCENT)
    rect(s, 0, 0, Inches(0.18), Inches(1.15), ACCENT)
    tb, tf = textbox(s, Inches(0.55), Inches(0.12), Inches(11.6), Inches(0.95), anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        add_par(tf, kicker.upper(), size=11.5, color=LIGHT, bold=True, first=True, space_after=2)
        add_par(tf, title, size=20, color=WHITE, bold=True)
    else:
        add_par(tf, title, size=23, color=WHITE, bold=True, first=True)
    if idx is not None:
        tbn, tfn = textbox(s, Inches(12.4), 0, Inches(0.9), Inches(1.15), anchor=MSO_ANCHOR.MIDDLE)
        add_par(tfn, idx, size=30, color=ACCENT, bold=True, align=PP_ALIGN.CENTER, first=True)


def footer(s, n):
    tb, tf = textbox(s, Inches(0.45), Inches(7.07), Inches(9), Inches(0.35))
    add_par(tf, "PLD/FT  |  Tipologias em Adquirencia  |  Documento confidencial - uso interno",
            size=8.5, color=RGBColor(0x90, 0x9C, 0xAC), first=True)
    tb2, tf2 = textbox(s, Inches(12.2), Inches(7.07), Inches(0.9), Inches(0.35))
    add_par(tf2, str(n), size=9, color=SECONDARY, bold=True, align=PP_ALIGN.RIGHT, first=True)


def newpage(kicker, title, idx=None):
    _pageno["n"] += 1
    s = slide()
    rect(s, 0, 0, SW, SH, WHITE)
    page_header(s, kicker, title, idx)
    footer(s, _pageno["n"])
    return s


def card(s, x, y, w, h, header, color=PRIMARY, header_color=WHITE, body_fill=ICE, hsize=12):
    rect(s, x, y, w, h, body_fill, line=LIGHT, line_w=Pt(0.75))
    hh = Inches(0.40)
    rect(s, x, y, w, hh, color)
    tbh, tfh = textbox(s, x + Inches(0.12), y, w - Inches(0.24), hh, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tfh, header, size=hsize, color=header_color, bold=True, first=True)
    tb, tf = textbox(s, x + Inches(0.14), y + hh + Inches(0.05), w - Inches(0.28), h - hh - Inches(0.12))
    return tf


def chip(s, x, y, w, text, color=SECONDARY, tcolor=WHITE, h=Inches(0.34), size=10.5):
    shp = rect(s, x, y, w, h, color, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = tcolor
    return shp


def prio_color(p):
    return {"Critica": RED, "Alta": AMBER, "Media": SECONDARY, "Baixa": GREEN}.get(p, GRAYTXT)


def style_table(table, body_size=9.5, header_size=10, col_aligns=None):
    tbl = table._tbl
    for tblPr in tbl.iter(qn('a:tblPr')):
        tblPr.set('firstRow', '1'); tblPr.set('bandRow', '0')
    for r in range(len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.07); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if (r % 2 == 1) else ICE
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.0
                if col_aligns and r > 0:
                    p.alignment = col_aligns[c]
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(header_size if r == 0 else body_size)
                    run.font.bold = (r == 0)
                    run.font.color.rgb = WHITE if r == 0 else GRAYTXT


def set_cell(table, r, c, text, color=None, bold=None, align=None):
    cell = table.cell(r, c)
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    if align is not None:
        p.alignment = align
    for run in p.runs:
        if color is not None:
            run.font.color.rgb = color
        if bold is not None:
            run.font.bold = bold


# ============================================================================
# CAPA
# ============================================================================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, SW, Inches(0.28), ACCENT)
rect(s, 0, Inches(5.55), SW, Inches(0.10), ACCENT)
rect(s, Inches(10.7), 0, Inches(2.63), SH, PRIMARY)
rect(s, Inches(10.4), 0, Inches(0.10), SH, ACCENT)
for i in range(11):
    hh = Inches(0.18 + (i % 5) * 0.12)
    rect(s, Inches(10.95), Inches(0.9) + Emu(int(i * 460000)), hh, Inches(0.10),
         ACCENT if i % 2 else SECONDARY)
tb, tf = textbox(s, Inches(0.7), Inches(1.25), Inches(9.4), Inches(0.6))
add_par(tf, "APRESENTACAO EXECUTIVA  |  VOLUME II  |  COMPLIANCE - RISCOS - PLD/FT",
        size=13, color=ACCENT, bold=True, first=True)
tb, tf = textbox(s, Inches(0.7), Inches(2.0), Inches(9.4), Inches(2.7))
add_par(tf, "Tipologias Especificas de Lavagem de Dinheiro e Financiamento do "
            "Terrorismo na Adquirencia",
        size=31, color=WHITE, bold=True, first=True, line_spacing=1.05, space_after=8)
add_par(tf, "Descricao, Indicadores de Alerta, Impacto, Controles e Monitoramento",
        size=18, color=ACCENT, bold=True, line_spacing=1.05)
tb, tf = textbox(s, Inches(0.7), Inches(5.9), Inches(9.4), Inches(1.3))
add_par(tf, "15 tipologias mapeadas para merchant acquirers, subcredenciadores, "
            "facilitadores de pagamento e marketplaces no Brasil",
        size=13, color=LIGHT, first=True, space_after=6, line_spacing=1.1)
add_par(tf, "Documento confidencial - uso interno  |  Comites Executivos e Auditoria",
        size=10, color=RGBColor(0x9A, 0xB4, 0xCE))


# ============================================================================
# AGENDA (15 tipologias)
# ============================================================================
tipos = [
    ("T1",  "Laranjas e Empresas de Fachada/Ficticias", "Critica"),
    ("T2",  "Saque Disfarcado no Cartao de Credito", "Alta"),
    ("T3",  "BF Oculto em PayFac, Marketplace e Subcredenciador", "Critica"),
    ("T4",  "Desvio de Recursos Publicos (CPGF / Sem Licitacao)", "Alta"),
    ("T5",  "Financiamento ao Terrorismo, Extremismo e Crime Org.", "Critica"),
    ("T6",  "Fraudes Internas", "Alta"),
    ("T7",  "Desvio de Recursos via ONGs", "Alta"),
    ("T8",  "Superfaturamento de Notas e Contratos de Licitacao", "Alta"),
    ("T9",  "Agiotagem e Autofinanciamento", "Alta"),
    ("T10", "Fronteiras e Ramos de Atividade de Risco", "Alta"),
    ("T11", "Comercializacao de Vouchers e Gift Cards", "Media"),
    ("T12", "Pagamento de Boletos por Terceiros", "Alta"),
    ("T13", "Jogos de Azar e Apostas", "Critica"),
    ("T14", "Estruturas Societarias Complexas (BF Oculto)", "Alta"),
    ("T15", "Contas de Passagem (Pass-through)", "Alta"),
]

s = newpage("Roteiro", "As 15 Tipologias Mapeadas para a Adquirencia", "00")
tb, tf = textbox(s, Inches(0.5), Inches(1.38), Inches(12.3), Inches(0.45))
add_par(tf, "Cada tipologia e detalhada em um slide com descricao, indicadores de alerta, "
            "impacto na adquirente, controles e regras de monitoramento:",
        size=12, color=GRAYTXT, first=True)
cw = Inches(3.97); chh = Inches(0.92); gx = Inches(0.16); gy = Inches(0.13)
sx = Inches(0.5); sy = Inches(1.95)
for i, (code, name, prio) in enumerate(tipos):
    r = i // 3; c = i % 3
    x = sx + c * (cw + gx); y = sy + r * (chh + gy)
    rect(s, x, y, cw, chh, WHITE, line=LIGHT, line_w=Pt(1))
    rect(s, x, y, Inches(0.78), chh, PRIMARY)
    tbc, tcc = textbox(s, x, y, Inches(0.78), chh, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tcc, code, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tcf = textbox(s, x + Inches(0.88), y + Inches(0.06), cw - Inches(0.98), chh - Inches(0.12),
                      anchor=MSO_ANCHOR.MIDDLE)
    add_par(tcf, name, size=10, color=PRIMARY, bold=True, first=True, line_spacing=0.98)
    rect(s, x + cw - Inches(0.16), y, Inches(0.16), chh, prio_color(prio))


# ============================================================================
# Funcao do slide de tipologia
# ============================================================================
def tipologia_slide(code, nome, idx, prio, descricao, indicadores, impacto, controles, regras):
    s = newpage("Tipologia " + code, nome, idx)
    # selo de prioridade
    chip(s, Inches(11.05), Inches(1.27), Inches(1.55),
         "PRIORIDADE: " + prio.upper(), color=prio_color(prio), size=9)

    # Descricao - faixa superior
    dy = Inches(1.45)
    tf = card(s, Inches(0.45), dy, Inches(12.43), Inches(1.12), "Descricao da tipologia", color=NAVY)
    add_par(tf, descricao, size=10.3, color=GRAYTXT, first=True, line_spacing=1.04)

    # Grade 2x2
    gy0 = Inches(2.72)
    cw = Inches(6.12); ch = Inches(1.98); gx = Inches(0.19); gyy = Inches(0.13)
    lx = Inches(0.45); rx = lx + cw + gx

    def fill(tf, items, mcolor=None):
        for it in items:
            bullet(tf, it, size=9.3, first=(it == items[0]), space_after=2, mcolor=mcolor)

    tf = card(s, lx, gy0, cw, ch, "Indicadores de alerta (red flags)", color=RED,
              body_fill=RGBColor(0xFB, 0xEC, 0xEA))
    fill(tf, indicadores, mcolor=RED)

    tf = card(s, rx, gy0, cw, ch, "Como a adquirente pode ser impactada", color=PRIMARY)
    fill(tf, impacto)

    gy1 = gy0 + ch + gyy
    tf = card(s, lx, gy1, cw, ch, "Controles preventivos e detectivos", color=SECONDARY)
    fill(tf, controles)

    tf = card(s, rx, gy1, cw, ch, "Regras de monitoramento sugeridas", color=ACCENT)
    fill(tf, regras)
    return s


# ---- T1 ----
tipologia_slide(
    "T1", "Uso de Laranjas, Empresas de Fachada e Ficticias", "01", "Critica",
    "Utilizacao de interpostas pessoas (laranjas) e de pessoas juridicas sem substancia economica real "
    "(empresas de fachada ou ficticias) para credenciar estabelecimentos comerciais e movimentar recursos "
    "ilicitos, ocultando o verdadeiro controlador e a origem dos valores.",
    ["Socio sem capacidade economica compativel com o volume processado.",
     "Endereco inexistente, virtual ou compartilhado por varios EC.",
     "EC recem-aberto com faturamento elevado e sem presenca verificavel.",
     "Conta de liquidacao em nome de terceiro diverso do titular do EC.",
     "Mesmo socio/BF vinculado a multiplos EC sem relacao operacional."],
    ["Credenciamento de EC inexistente usado para colocacao de recursos.",
     "Rede de pagamentos utilizada como veiculo de lavagem.",
     "Exposicao reputacional e sancionatoria perante BCB/COAF.",
     "Aumento de chargeback e perdas associadas a EC de fachada."],
    ["KYB robusto e identificacao do beneficiario final.",
     "Validacao cadastral, biometria e prova de vida no onboarding.",
     "Geolocalizacao e visita ao estabelecimento para EC de risco.",
     "Screening de socios, BF e midia adversa."],
    ["Conta de liquidacao com CPF/CNPJ diferente do titular do EC.",
     "N+ EC com mesmo socio, endereco, telefone ou IP.",
     "EC ativo ha menos de 90 dias com volume acima de R$ X.",
     "Socio com flag de laranja em base interna/externa."],
)

# ---- T2 ----
tipologia_slide(
    "T2", "Desvio de Recursos via Emprestimo Irregular no Cartao de Credito", "02", "Alta",
    "Uso do estabelecimento para simular vendas e converter o limite do cartao de credito do portador em "
    "dinheiro (saque disfarcado / 'emprestimo' informal), normalmente mediante agio. Configura credito "
    "irregular, autofinanciamento e lavagem por meio da maquininha, sem entrega de mercadoria.",
    ["Transacoes de valores redondos sem lastro de mercadoria/servico.",
     "Mesmo portador recorrente em um ou poucos EC.",
     "Ticket elevado e incompativel com o porte/MCC do EC.",
     "Excesso de parcelamentos e estornos atipicos.",
     "EC que anuncia 'saque no cartao' ou 'dinheiro na hora'."],
    ["Rede usada para concessao de credito irregular a juros.",
     "Aumento de inadimplencia, contestacao e chargeback.",
     "Risco legal (agiotagem) e descumprimento de regras de bandeira.",
     "Distorcao do perfil transacional e da carteira de EC."],
    ["Monitoramento de padrao de saque disfarcado (ticket x MCC).",
     "Regras das bandeiras contra cash advance disfarcado.",
     "Due diligence reforcada e limites por perfil de EC.",
     "Analise de recorrencia portador-estabelecimento."],
    ["Alta proporcao de transacoes em valores redondos no EC.",
     "Recorrencia anomala do mesmo portador no mesmo EC.",
     "Ticket medio > X em MCC incompativel com a faixa.",
     "Pico de parcelamento + estorno na mesma janela."],
)

# ---- T3 ----
tipologia_slide(
    "T3", "Ocultacao de BF em PayFac Internacional, Marketplace e Subcredenciador", "03", "Critica",
    "Estruturas de facilitadores de pagamento internacionais, marketplaces e subcredenciadores agregam os "
    "estabelecimentos finais sob um unico relacionamento, reduzindo a visibilidade do EC e ocultando o "
    "beneficiario final por meio de nested merchants e fluxos transfronteiricos.",
    ["Recusa do parceiro em fornecer dados do EC final (look-through).",
     "Sellers onboarded com diligencia minima e crescimento explosivo.",
     "Alta dispersao de MCC e geografias sob um unico CNPJ/contrato.",
     "Fluxo cross-border elevado e estruturas em offshore.",
     "Chargeback e fraude concentrados em poucos sellers/EC."],
    ["EC final invisivel a adquirente (nested merchants).",
     "Rede usada para transaction laundering por terceiros.",
     "Exposicao a sancoes e a jurisdicoes de risco no cross-border.",
     "Dificuldade de identificar o beneficiario final exigido por norma."],
    ["Due diligence reforcada do subcredenciador/PayFac/marketplace.",
     "Exigencia contratual de KYC de seller e direito de look-through.",
     "Avaliacao do programa de PLD/FT e do KYB do parceiro.",
     "Identificacao do BF do EC final e screening da cadeia."],
    ["Seller individual com volume/chargeback acima de limiar.",
     "Fluxo cross-border > X% do volume do parceiro.",
     "Sellers compartilhando IP, conta de liquidacao ou BF.",
     "Crescimento de sellers > Y% sem reforco de diligencia."],
)

# ---- T4 ----
tipologia_slide(
    "T4", "Desvio de Recursos Publicos (Cartao Corporativo / Sem Licitacao)", "04", "Alta",
    "Uso de cartao de pagamento do governo (CPGF) ou de empresas contratadas sem licitacao (dispensa/"
    "fracionamento) para desviar recursos publicos por meio de estabelecimentos, frequentemente de fachada, "
    "vinculados a agentes publicos ou a fornecedores sem capacidade operacional.",
    ["EC recebendo de cartoes corporativos publicos com saques/valores atipicos.",
     "Fornecedor sem estrutura compativel com os contratos firmados.",
     "Concentracao de pagamentos publicos em poucos EC.",
     "Valores fracionados logo abaixo do limite de dispensa de licitacao.",
     "Vinculo do EC/BF a agente publico ou a PEP."],
    ["EC de fachada inserido na rede para captar recursos publicos.",
     "Associacao a corrupcao e desvio, com risco reputacional grave.",
     "Exposicao a investigacoes de orgaos de controle.",
     "Comunicacao obrigatoria de operacoes suspeitas ao COAF."],
    ["KYB e avaliacao de capacidade do fornecedor.",
     "Identificacao de PEP e de vinculos com entes publicos.",
     "Monitoramento especifico de cartoes corporativos/CPGF.",
     "Midia adversa e screening reforcado."],
    ["Padrao de saque ou de valores atipicos em CPGF.",
     "Pagamentos recorrentes logo abaixo do limite de dispensa.",
     "EC vinculado a PEP ou a ente publico contratante.",
     "Concentracao de recebiveis de origem publica em poucos EC."],
)

# ---- T5 ----
tipologia_slide(
    "T5", "Financiamento de Grupos Extremistas, Criminosos e Terroristas (FT)", "05", "Critica",
    "Captacao, movimentacao e disponibilizacao de recursos a individuos e organizacoes terroristas, "
    "extremistas ou criminosas por meio de cartoes, doacoes e EC/ONG de fachada. Caracteriza-se por baixos "
    "valores, captacao difusa e nexo com listas de sancoes, exigindo bloqueio imediato de ativos.",
    ["Match de EC, socio, BF ou contraparte com listas de sancoes.",
     "Doacoes pulverizadas de muitas origens para um mesmo EC/ONG.",
     "Nexo geografico com zonas de conflito ou jurisdicoes de risco.",
     "ONG/entidade sem transparencia de finalidade e destino.",
     "Vendas simbolicas e baixos valores recorrentes."],
    ["Rede usada para arrecadar e repassar recursos a grupos alvo.",
     "Obrigacao de indisponibilidade imediata de ativos (Lei 13.810/2019).",
     "Risco penal, reputacional e de sancao regulatoria.",
     "Comunicacao ao COAF independentemente do valor."],
    ["Screening de sancoes (CSNU/OFAC/UE/UK) continuo.",
     "Identificacao de BF e da finalidade de ONG/NPO.",
     "Due diligence reforcada para entidades de risco.",
     "Processo de bloqueio imediato e reporte."],
    ["Match em lista de sancoes/FT (qualquer valor).",
     "N+ doacoes pulverizadas para EC/ONG em janela curta.",
     "Transacoes com nexo a jurisdicao de conflito.",
     "Beneficiario/keyword em watchlist de FT."],
)

# ---- T6 ----
tipologia_slide(
    "T6", "Ocultacao de Recursos Associada a Fraudes Internas", "06", "Alta",
    "Desvio de recursos por colaboradores, parceiros ou prestadores, com ocultacao por meio de "
    "credenciamentos manipulados, overrides indevidos de regras de risco, ajustes de liquidacao e EC "
    "vinculados a insiders, comprometendo a integridade dos controles da adquirente.",
    ["EC com BF, conta ou contato vinculado a colaborador/parceiro.",
     "Overrides manuais frequentes de regras de risco/onboarding.",
     "Alteracoes de conta de liquidacao realizadas por insiders.",
     "Aprovacoes fora de alcada ou sem trilha de auditoria.",
     "Acessos e ajustes em horarios e padroes atipicos."],
    ["Perdas financeiras diretas e fraude transacional.",
     "Comprometimento da efetividade dos controles de PLD.",
     "Risco reputacional e de conformidade.",
     "Mascaramento de outras tipologias por agentes internos."],
    ["Segregacao de funcoes e principio dos quatro olhos.",
     "Trilha de auditoria completa e imutavel.",
     "KYC de colaborador e gestao de conflito de interesse.",
     "Revisao periodica de overrides e excecoes."],
    ["EC com BF/conta associada a colaborador (link analysis).",
     "Override de regra de risco acima do baseline do operador.",
     "Mudanca de conta de liquidacao sem dupla aprovacao.",
     "Concentracao de excecoes por usuario/area."],
)

# ---- T7 ----
tipologia_slide(
    "T7", "Desvio de Recursos por meio de ONGs", "07", "Alta",
    "Organizacoes da sociedade civil, associacoes e fundacoes utilizadas para captar doacoes via cartao e "
    "desviar ou lavar recursos, valendo-se da baixa transparencia de finalidade e do tratamento de risco "
    "diferenciado conferido ao terceiro setor.",
    ["ONG com volume de doacoes incompativel com o porte/atuacao.",
     "Doacoes pulverizadas e fluxo cross-border atipico.",
     "Dirigentes/BF com flags, PEP ou midia adversa.",
     "Finalidade generica e ausencia de prestacao de contas.",
     "Picos de captacao sem campanha ou evento correspondente."],
    ["Rede usada para lavagem ou FT via terceiro setor.",
     "Exposicao reputacional por associacao a entidade desviante.",
     "Dificuldade de identificar o real destino dos recursos.",
     "Comunicacao de operacoes suspeitas ao COAF."],
    ["KYB e due diligence reforcada de NPO.",
     "Identificacao de dirigentes e do beneficiario final.",
     "Screening de sancoes/PEP e analise de finalidade.",
     "Acompanhamento de transparencia e prestacao de contas."],
    ["Volume de doacoes acima do baseline da entidade.",
     "Pulverizacao de doacoes em janela curta.",
     "Nexo cross-border com jurisdicao de risco.",
     "Dirigente/BF com flag de sancao ou midia adversa."],
)

# ---- T8 ----
tipologia_slide(
    "T8", "Superfaturamento de Notas e Contratos de Licitacao", "08", "Alta",
    "Emissao de notas fiscais e contratos com valores inflados (over-invoicing) para justificar a "
    "movimentacao e desviar recursos, modalidade de trade-based money laundering que pode transitar pela "
    "rede de adquirencia via cartao ou boleto, especialmente em contratacoes publicas.",
    ["Ticket muito acima do preco de mercado/referencia.",
     "Notas e contratos sem lastro economico verificavel.",
     "Concentracao de pagamentos entre o mesmo par comprador-EC.",
     "Precos divergentes de tabelas de referencia (ex.: publicas).",
     "Aditivos e reajustes sucessivos sem justificativa."],
    ["Pagamentos inflados transitando pela rede.",
     "Lavagem por superfaturamento (TBML) via adquirencia.",
     "Exposicao a investigacoes de fraude em licitacoes.",
     "Risco reputacional e de conformidade."],
    ["Validacao de lastro e de nota fiscal.",
     "Monitoramento de preco vs. referencia de mercado.",
     "KYB e DDR para contratos publicos e grandes fornecedores.",
     "Analise de concentracao comprador-fornecedor."],
    ["Desvio de preco do EC vs. referencia/peer (> X%).",
     "Concentracao do par comprador-EC acima de limiar.",
     "Ticket fora da faixa esperada do MCC.",
     "Sequencia de aditivos elevando o valor pago."],
)

# ---- T9 ----
tipologia_slide(
    "T9", "Agiotagem e Autofinanciamento", "09", "Alta",
    "Uso da maquininha para conceder emprestimos informais a juros (agiotagem) ou para o proprio "
    "estabelecimento gerar liquidez de forma circular (autofinanciamento), por meio de autocompras com "
    "cartoes do socio/BF ou de laranjas, simulando faturamento.",
    ["Autocompras com cartoes do proprio socio/BF do EC.",
     "Transacoes circulares entre EC e portadores relacionados.",
     "Recorrencia anomala do mesmo portador no EC.",
     "Tickets padronizados e faturamento sem operacao real.",
     "Picos de venda seguidos de antecipacao integral."],
    ["Concessao de credito ilegal e liquidez artificial via rede.",
     "Aumento de chargeback e de risco de credito.",
     "Risco legal e descumprimento de regras de bandeira.",
     "Distorcao do perfil de risco do EC."],
    ["Link analysis entre BF do EC e portadores.",
     "Regras de deteccao de autocompra e circularidade.",
     "Limites de processamento e reserva por risco.",
     "Due diligence reforcada e analise de antecipacao."],
    ["Cartao do portador = socio/BF do EC.",
     "Padrao circular de transacoes (mesmas partes).",
     "Recorrencia portador-EC acima do baseline.",
     "Pico de vendas + antecipacao integral em janela curta."],
)

# ---- T10 ----
tipologia_slide(
    "T10", "Ocultacao em Regiao de Fronteira ou Ramos de Atividade de Risco", "10", "Alta",
    "Estabelecimentos localizados em zonas de fronteira ou atuantes em ramos de alto risco (cambio informal, "
    "combustiveis, joias e metais, postos, casas de cambio, importados) utilizados para colocacao e ocultacao "
    "de recursos, muitas vezes ligados a contrabando, descaminho e cambio paralelo.",
    ["EC em faixa de fronteira com volume e ticket atipicos.",
     "MCC de ramo de alto risco (cambio, joias, combustiveis).",
     "Fluxo cross-border elevado e sazonalidade incompativel.",
     "Pagamentos em especie/cartao desproporcionais a praca.",
     "Estrutura de EC vinculada a importacao/exportacao opaca."],
    ["Rede usada para colocacao em pontos de maior vulnerabilidade.",
     "Exposicao a contrabando, descaminho e cambio paralelo.",
     "Risco geografico e setorial elevado.",
     "Risco reputacional e de conformidade."],
    ["Triagem geografica e matriz de risco por regiao/MCC.",
     "Due diligence reforcada para ramos de risco.",
     "Geolocalizacao e visita ao estabelecimento.",
     "Monitoramento setorial e de cross-border."],
    ["EC em regiao de fronteira com volume > X (outlier regional).",
     "MCC de risco combinado a fluxo cross-border.",
     "Ticket/volume fora do baseline da praca.",
     "Sazonalidade incompativel com o ramo declarado."],
)

# ---- T11 ----
tipologia_slide(
    "T11", "Desvio de Recursos por meio de Comercializacao de Vouchers", "11", "Media",
    "Compra e revenda de vouchers, gift cards, recargas e creditos pre-pagos para converter e movimentar "
    "recursos ilicitos com baixa rastreabilidade, criando uma camada de layering de dificil reconstituicao "
    "da trilha financeira.",
    ["Compras elevadas e recorrentes de vouchers/gift cards.",
     "Revenda e resgate por terceiros sem relacao aparente.",
     "Tickets uniformes e em volumes atipicos.",
     "EC de recarga/voucher com faturamento incompativel.",
     "Padrao compra-resgate concentrado em curto intervalo."],
    ["Rede usada para layering por meio de instrumentos pre-pagos.",
     "Baixa rastreabilidade do destino final dos recursos.",
     "Exposicao a fraude e a uso indevido de creditos.",
     "Risco reputacional e de conformidade."],
    ["Monitoramento de MCC de voucher/gift card/recarga.",
     "Limites de compra e KYC reforcado.",
     "Analise de padrao compra-resgate e de contrapartes.",
     "Due diligence de EC emissor/revendedor."],
    ["Volume de vouchers acima do baseline do EC.",
     "Padrao de compra e resgate por terceiros distintos.",
     "Concentracao de tickets uniformes em janela curta.",
     "EC de recarga com crescimento atipico de volume."],
)

# ---- T12 ----
tipologia_slide(
    "T12", "Desvio de Recursos por Pagamento de Boletos por Terceiros", "12", "Alta",
    "Uso de estabelecimentos ou servicos para pagamento de boletos/contas de terceiros, criando contas de "
    "passagem que ofuscam a origem e o destino dos recursos e fragmentam a trilha financeira (layering).",
    ["Pagamentos de boletos sem relacao com a atividade do EC.",
     "Grande numero de beneficiarios distintos a partir de um EC.",
     "Valores fracionados e terceiros pagadores recorrentes.",
     "Origem dos recursos diversa do beneficiario final.",
     "Volume de pagamentos de contas incompativel com o porte."],
    ["Rede usada como conta de passagem (layering).",
     "Ofuscacao da origem e destino dos recursos.",
     "Dificuldade de reconstituir a trilha financeira.",
     "Comunicacao de operacoes suspeitas ao COAF."],
    ["Monitoramento de pagamento de contas/boletos por terceiros.",
     "KYC do pagador e do beneficiario quando aplicavel.",
     "Limites e regras por finalidade e perfil.",
     "Analise de rede de pagadores e beneficiarios."],
    ["N+ boletos pagos a beneficiarios distintos por um EC.",
     "Origem do recurso diferente do beneficiario.",
     "Fracionamento de pagamentos em janela curta.",
     "Recorrencia de terceiros pagadores sem vinculo."],
)

# ---- T13 ----
tipologia_slide(
    "T13", "Movimentacao de Recursos via Jogos de Azar ou Apostas", "13", "Critica",
    "Movimentacao de recursos ilicitos por meio de operadores de jogos de azar e apostas (legais e "
    "clandestinos), inclusive com MCC mascarado, gateways e fluxos cross-border, frequentemente associada a "
    "transaction laundering e a operadores sem licenca.",
    ["Soft descriptor ou URL com termos de aposta/jogo/'bet'.",
     "Payouts (pagamentos a apostadores) atipicos e bidirecionais.",
     "MCC incompativel com indicios de aposta/jogo.",
     "Fluxo cross-border para operadores e gateways.",
     "Operador sem licenca/autorizacao aplicavel."],
    ["Rede usada por operadores de apostas ilegais.",
     "Transaction laundering e ocultacao de receitas de jogo.",
     "Sancao das bandeiras (programas BRAM/VIRP).",
     "Risco reputacional e regulatorio elevado."],
    ["KYB setorial e checagem de licenca/autorizacao.",
     "Deteccao de keyword em descriptor/site/URL.",
     "Monitoramento de payout e fluxo bidirecional.",
     "Due diligence reforcada e politica por vertical."],
    ["Keyword de aposta/jogo em descriptor ou URL.",
     "Razao de payout/credito atipica para o MCC.",
     "Fluxo a operador sem licenca ou cross-border de risco.",
     "MCC declarado divergente do comportamento transacional."],
)

# ---- T14 ----
tipologia_slide(
    "T14", "Ocultacao de Beneficiarios Finais por Estruturas Societarias Complexas", "14", "Alta",
    "Uso de camadas de holdings, sociedades de proposito especifico, socios pessoa juridica e veiculos em "
    "offshore para ocultar o beneficiario final do estabelecimento, dificultando a identificacao exigida "
    "pela regulacao e o screening da cadeia.",
    ["Cadeia societaria com multiplas camadas e/ou offshore.",
     "Pessoa juridica controlando pessoa juridica em sequencia.",
     "Beneficiario final nao identificavel ou inconclusivo.",
     "Presenca de nominee directors/shareholders.",
     "Jurisdicoes de sigilo e baixa transparencia na estrutura."],
    ["Beneficiario final oculto, com risco de nao conformidade.",
     "Dificuldade de screening e de avaliacao de risco do EC.",
     "Veiculo para outras tipologias (fachada, FT, sancoes).",
     "Risco reputacional e regulatorio."],
    ["Identificacao do BF ate a pessoa natural (look-through).",
     "Analise e mapeamento da cadeia societaria.",
     "Due diligence reforcada para estruturas opacas.",
     "Screening de toda a cadeia de controle."],
    ["Cadeia societaria com mais de N camadas.",
     "Presenca de offshore/jurisdicao de sigilo na estrutura.",
     "BF nao conclusivo apos diligencia padrao.",
     "Reuso de mesma estrutura/agentes em multiplos EC."],
)

# ---- T15 ----
tipologia_slide(
    "T15", "Uso de Contas de Passagem (Pass-through)", "15", "Alta",
    "Utilizacao de contas ou estabelecimentos de passagem para receber e repassar rapidamente recursos, sem "
    "retencao nem proposito economico real, fragmentando a trilha financeira e dificultando o rastreamento "
    "da origem e do destino (layering).",
    ["Entradas e saidas casadas em curtissimo intervalo.",
     "Saldo recorrentemente proximo de zero (turnover alto).",
     "Ausencia de atividade comercial real no EC.",
     "Multiplas contrapartes de origem e destino (many-to-many).",
     "Volume incompativel com o porte e a finalidade declarada."],
    ["Rede usada como camada de passagem para layering.",
     "Dificuldade de rastrear origem e destino dos recursos.",
     "Mascaramento de outras tipologias subjacentes.",
     "Comunicacao de operacoes suspeitas ao COAF."],
    ["Monitoramento de fluxo entra-sai e proposito economico.",
     "Analise de contrapartes e de rede de transacoes.",
     "KYC/KYB e validacao de atividade real.",
     "Holds e reservas por regra de risco."],
    ["Razao saida/entrada proxima de 1 em janela curta.",
     "Turnover elevado com saldo medio baixo.",
     "Padrao many-to-many de contrapartes.",
     "Ausencia de retencao e de margem operacional."],
)


# ============================================================================
# QUADRO-RESUMO / MATRIZ DE PRIORIZACAO
# ============================================================================
s = newpage("Sintese", "Matriz de Priorizacao e Vetor Principal de Controle", "16")
cols = ["Cod.", "Tipologia", "Vetor principal de deteccao", "Prioridade"]
widths = [Inches(0.8), Inches(4.4), Inches(5.6), Inches(1.63)]
matrix = [
    ("T1", "Laranjas e empresas de fachada", "KYB + beneficiario final + conta de liquidacao x titular", "Critica"),
    ("T2", "Saque disfarcado no cartao de credito", "Padrao de valores redondos + recorrencia portador-EC", "Alta"),
    ("T3", "BF oculto em PayFac/marketplace/subcred.", "Look-through + KYC de seller + monitoramento por parceiro", "Critica"),
    ("T4", "Desvio de recursos publicos", "PEP + monitoramento de CPGF + limites de dispensa", "Alta"),
    ("T5", "Financiamento ao terrorismo/extremismo", "Screening de sancoes continuo + bloqueio imediato", "Critica"),
    ("T6", "Fraudes internas", "Segregacao de funcoes + trilha + revisao de overrides", "Alta"),
    ("T7", "Desvio de recursos via ONGs", "DDR de NPO + BF + analise de finalidade/doacoes", "Alta"),
    ("T8", "Superfaturamento de notas/licitacoes", "Preco vs. referencia + lastro de NF + concentracao", "Alta"),
    ("T9", "Agiotagem e autofinanciamento", "Link analysis BF-portador + autocompra/circularidade", "Alta"),
    ("T10", "Fronteiras e ramos de risco", "Triagem geografica/setorial + cross-border", "Alta"),
    ("T11", "Comercializacao de vouchers", "MCC de voucher + padrao compra-resgate + limites", "Media"),
    ("T12", "Pagamento de boletos por terceiros", "Rede de pagadores/beneficiarios + fracionamento", "Alta"),
    ("T13", "Jogos de azar e apostas", "Keyword descriptor/URL + payout + licenca", "Critica"),
    ("T14", "Estruturas societarias complexas", "Look-through da cadeia + screening + DDR", "Alta"),
    ("T15", "Contas de passagem", "Fluxo entra-sai + turnover/saldo + many-to-many", "Alta"),
]
nrows = len(matrix) + 1
gtbl = s.shapes.add_table(nrows, 4, Inches(0.45), Inches(1.4), sum(widths, Emu(0)), Inches(5.5))
table = gtbl.table
for i, w in enumerate(widths):
    table.columns[i].width = w
for c, h in enumerate(cols):
    table.cell(0, c).text = h
for r, row in enumerate(matrix, start=1):
    for c, val in enumerate(row):
        table.cell(r, c).text = val
aligns = [PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER]
style_table(table, body_size=9.3, header_size=10, col_aligns=aligns)
for r, row in enumerate(matrix, start=1):
    set_cell(table, r, 0, row[0], color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    set_cell(table, r, 1, row[1], color=GRAYTXT, bold=True, align=PP_ALIGN.LEFT)
    set_cell(table, r, 3, row[3], color=prio_color(row[3]), bold=True, align=PP_ALIGN.CENTER)
table.rows[0].height = Inches(0.34)


# ============================================================================
# ENCERRAMENTO
# ============================================================================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(3.2), SW, Inches(0.08), ACCENT)
rect(s, Inches(10.7), 0, Inches(2.63), SH, PRIMARY)
rect(s, Inches(10.4), 0, Inches(0.10), SH, ACCENT)
tb, tf = textbox(s, Inches(0.8), Inches(2.4), Inches(9.3), Inches(1.2))
add_par(tf, "Mensagem-chave", size=14, color=ACCENT, bold=True, first=True, space_after=6)
add_par(tf, "As 15 tipologias compartilham vetores comuns de controle: identificacao do "
            "beneficiario final, KYC/KYB, screening de sancoes e monitoramento transacional "
            "calibrado por risco.",
        size=18, color=WHITE, bold=True, line_spacing=1.15)
tb, tf = textbox(s, Inches(0.8), Inches(4.6), Inches(9.3), Inches(1.6))
add_par(tf, "A priorizacao por risco (critica/alta/media) orienta a alocacao de esforco "
            "da esteira de prevencao, com governanca de tuning, trilha de auditoria e "
            "comunicacao tempestiva ao COAF, em linha com a Circular BCB 3.978/2020 e as "
            "recomendacoes do GAFI/FATF.",
        size=12.5, color=LIGHT, first=True, line_spacing=1.2)
tb, tf = textbox(s, Inches(0.8), Inches(6.6), Inches(9.3), Inches(0.5))
add_par(tf, "Documento confidencial - uso interno  |  Compliance, Riscos e PLD/FT",
        size=10, color=RGBColor(0x9A, 0xB4, 0xCE), first=True)


out = "/home/user/localiza_leilao/Apresentacao_PLD_Tipologias_Adquirencia_Vol2.pptx"
prs.save(out)
print("Slides:", len(prs.slides._sldIdLst))
print("Arquivo salvo em:", out)
