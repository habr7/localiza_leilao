#!/usr/bin/env python3
"""Gera um one-pager (slide único) sobre o diferencial do Barbacoa em São Paulo."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Paleta de cores ----
PRETO      = RGBColor(0x1A, 0x14, 0x10)   # fundo escuro / churrascaria
MARROM     = RGBColor(0x3B, 0x29, 0x1E)
DOURADO    = RGBColor(0xC4, 0x9A, 0x44)   # dourado premium
DOURADO_CL = RGBColor(0xE6, 0xC8, 0x7A)
BRANCO     = RGBColor(0xF7, 0xF2, 0xE9)
CINZA_CL   = RGBColor(0xBD, 0xB3, 0xA4)
VINHO      = RGBColor(0x6E, 0x1A, 0x1A)

# ---- Apresentação 16:9 ----
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout em branco


def add_rect(left, top, width, height, fill, line=None, line_w=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    return shp


def add_text(left, top, width, height, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    """lines: lista de dicts {text, size, color, bold, italic, space_after}"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if "space_after" in ln:
            p.space_after = Pt(ln["space_after"])
        if "space_before" in ln:
            p.space_before = Pt(ln["space_before"])
        if "line_spacing" in ln:
            p.line_spacing = ln["line_spacing"]
        r = p.add_run()
        r.text = ln["text"]
        f = r.font
        f.size = Pt(ln.get("size", 14))
        f.bold = ln.get("bold", False)
        f.italic = ln.get("italic", False)
        f.color.rgb = ln.get("color", BRANCO)
        f.name = ln.get("font", "Georgia")
    return tb


# ===== Fundo =====
add_rect(0, 0, SW, SH, PRETO)

# Barra lateral esquerda (faixa premium)
SIDE_W = Inches(4.3)
add_rect(0, 0, SIDE_W, SH, MARROM)
# linha dourada vertical de separação
add_rect(SIDE_W - Pt(3), 0, Pt(3), SH, DOURADO)

# ===== Cabeçalho na faixa lateral =====
add_text(Inches(0.5), Inches(0.6), SIDE_W - Inches(0.9), Inches(2.2), [
    {"text": "BARBACOA", "size": 46, "bold": True, "color": DOURADO,
     "font": "Georgia", "space_after": 2},
    {"text": "C H U R R A S C A R I A", "size": 13, "bold": True,
     "color": DOURADO_CL, "font": "Georgia", "space_after": 14},
    {"text": "Steakhouse premium", "size": 15, "italic": True,
     "color": BRANCO, "font": "Georgia", "space_after": 0},
    {"text": "São Paulo · desde 1990", "size": 13, "color": CINZA_CL,
     "font": "Georgia"},
])

# Bloco "Por que somos referência" na lateral
add_rect(Inches(0.5), Inches(3.05), SIDE_W - Inches(1.0), Pt(2), DOURADO)
add_text(Inches(0.5), Inches(3.25), SIDE_W - Inches(0.9), Inches(3.8), [
    {"text": "POR QUE SOMOS REFERÊNCIA", "size": 13, "bold": True,
     "color": DOURADO, "font": "Georgia", "space_after": 12},
    {"text": "★  Tradição de mais de 30 anos no alto",
     "size": 13.5, "color": BRANCO, "space_after": 2, "line_spacing": 1.0},
    {"text": "     padrão da gastronomia paulistana",
     "size": 13.5, "color": BRANCO, "space_after": 10},
    {"text": "★  Endereços nobres: Itaim Bibi,",
     "size": 13.5, "color": BRANCO, "space_after": 2, "line_spacing": 1.0},
    {"text": "     Jardins e Morumbi",
     "size": 13.5, "color": BRANCO, "space_after": 10},
    {"text": "★  Marca brasileira reconhecida",
     "size": 13.5, "color": BRANCO, "space_after": 2, "line_spacing": 1.0},
    {"text": "     internacionalmente (Tóquio)",
     "size": 13.5, "color": BRANCO, "space_after": 10},
    {"text": "★  Experiência completa: do corte",
     "size": 13.5, "color": BRANCO, "space_after": 2, "line_spacing": 1.0},
    {"text": "     nobre ao serviço impecável",
     "size": 13.5, "color": BRANCO},
])

# ===== Título da área principal =====
MAIN_L = SIDE_W + Inches(0.55)
MAIN_W = SW - MAIN_L - Inches(0.55)

add_text(MAIN_L, Inches(0.55), MAIN_W, Inches(1.0), [
    {"text": "O DIFERENCIAL BARBACOA", "size": 27, "bold": True,
     "color": DOURADO, "font": "Georgia", "space_after": 2},
    {"text": "A melhor experiência de churrascaria em São Paulo",
     "size": 14, "italic": True, "color": CINZA_CL, "font": "Georgia"},
])

# ===== Cards de diferenciais (2x2) =====
cards = [
    ("CARNES NOBRES & MATURAÇÃO",
     "Cortes premium selecionados e maturados com rigor, "
     "assados na brasa no ponto perfeito por mestres churrasqueiros."),
    ("RODÍZIO DE ALTO PADRÃO",
     "Variedade contínua de cortes servidos à mesa, no seu ritmo, "
     "com qualidade constante do primeiro ao último."),
    ("BUFFET GOURMET COMPLETO",
     "Ilha de saladas, frutos do mar, pratos quentes e queijos finos — "
     "uma curadoria que vai muito além do churrasco."),
    ("AMBIENTE & SERVIÇO PREMIUM",
     "Salão sofisticado e atendimento atencioso para negócios, "
     "celebrações e ocasiões especiais no coração de SP."),
]

card_w = (MAIN_W - Inches(0.4)) / 2
card_h = Inches(2.05)
gx, gy = Inches(0.4), Inches(0.35)
start_top = Inches(1.95)

for idx, (titulo, desc) in enumerate(cards):
    row, col = divmod(idx, 2)
    left = MAIN_L + col * (card_w + gx)
    top  = start_top + row * (card_h + gy)
    card = add_rect(left, top, card_w, card_h, MARROM,
                    line=DOURADO, line_w=Pt(1))
    # detalhe dourado no topo do card
    add_rect(left, top, card_w, Pt(4), DOURADO)
    add_text(left + Inches(0.22), top + Inches(0.25),
             card_w - Inches(0.44), card_h - Inches(0.4), [
        {"text": titulo, "size": 15, "bold": True, "color": DOURADO_CL,
         "font": "Georgia", "space_after": 8},
        {"text": desc, "size": 12.5, "color": BRANCO, "font": "Georgia",
         "line_spacing": 1.12},
    ])

# ===== Rodapé (faixa) =====
add_rect(MAIN_L, Inches(6.75), MAIN_W, Inches(0.5), VINHO)
add_text(MAIN_L + Inches(0.25), Inches(6.78), MAIN_W - Inches(0.5),
         Inches(0.44), [
    {"text": "Tradição, qualidade e sofisticação em cada experiência.",
     "size": 13, "bold": True, "italic": True, "color": DOURADO_CL,
     "font": "Georgia", "align": PP_ALIGN.CENTER}],
    anchor=MSO_ANCHOR.MIDDLE)

out = "/home/user/localiza_leilao/Barbacoa_OnePager_SaoPaulo.pptx"
prs.save(out)
print("Salvo em:", out)
